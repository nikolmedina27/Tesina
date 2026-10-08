"""Genera entregables/Ejemplo_importacion_OT-DEMO-001.xlsx: el formato único LLENO con datos DEMO.

Sirve para probar la importación y mostrar el avance por pieza y la curva S del proyecto de demostración.
Las fechas se calculan desde el inicio del proyecto OT-DEMO-001 en data/plataforma.db (o hace 24 días).
Todos los datos son inventados (DEMO). Uso: py scripts/generar_ejemplo_importacion.py
"""
import random
import sqlite3
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / 'scripts'))
from generar_formato_unico import main as generar_plantilla  # noqa: E402

SALIDA = RAIZ / 'entregables' / 'Ejemplo_importacion_OT-DEMO-001.xlsx'
OT = 'OT-DEMO-001'


def inicio_demo():
    try:
        con = sqlite3.connect(RAIZ / 'data' / 'plataforma.db')
        r = con.execute('SELECT inicio FROM proyecto_activo WHERE codigo=?', (OT,)).fetchone()
        if r:
            return date.fromisoformat(r[0])
    except sqlite3.Error:
        pass
    return date.today() - timedelta(days=24)


def laborable(d):
    while d.weekday() == 6:
        d += timedelta(days=1)
    return d


def escribir(ws, filas):
    """Escribe filas (dict título→valor) desde la fila 3, borrando el ejemplo y su marcador."""
    titulos = {str(ws.cell(2, j).value).strip(): j for j in range(1, ws.max_column + 1) if ws.cell(2, j).value}
    for j in range(1, ws.max_column + 2):
        c = ws.cell(3, j)
        if not (isinstance(c.value, str) and c.value.startswith('=')):
            c.value = None
    for i, f in enumerate(filas, start=3):
        for t, v in f.items():
            ws.cell(i, titulos[t], v)


def main():
    random.seed(11)
    hoy, ini = date.today(), inicio_demo()
    plantilla = generar_plantilla(SALIDA)
    wb = openpyxl.load_workbook(plantilla)

    escribir(wb['PROYECTO'], [{'Código OT': OT, 'Cliente': 'Cliente demo', 'Nombre del proyecto': 'Nave de almacenamiento (DEMO)',
                               'Tipo de estructura': 'Nave industrial / packing / cobertura', 'Toneladas metrado': 47,
                               'Fuente de tonelaje': 'Planos de fabricación', 'N° piezas': 0, '% planchas': 0.17, 'm² de pintura': 660,
                               'Metros de soldadura': 850, 'Incluye montaje': 'No', 'Presupuesto sin IGV': 150000, 'Moneda': 'USD',
                               'Día 1 contractual': ini.isoformat(), 'Fecha fin contractual': (ini + timedelta(days=62)).isoformat(),
                               'Resultado cotización': 'Ganada'}])

    tipos = [('COLUMNA', 'W8x31', 410, 24), ('VIGA', 'W10x22', 320, 26), ('TIJERAL', 'L3x3x1/4', 880, 12), ('CORREA', 'C6x8.2', 62, 120),
             ('ARRIOSTRE', 'L2x2x3/16', 24, 60), ('PLACA BASE', 'PL 3/4"', 34, 24), ('PLACA DE CONEXIÓN', 'PL 1/2"', 12, 96),
             ('TEMPLADOR', 'Ø5/8"', 8, 80), ('INSERTO', 'PL 1/2"', 40, 48)]
    contr = ['T&C FABRICACION', 'FRANCISCO TARRILLO', 'RFR METALICAS', 'ROGGER TORRES']
    piezas, n = [], 0
    for tipo, perfil, pu, cant_total in tipos:
        lotes = max(1, cant_total // 6)
        for k in range(lotes):
            n += 1
            cant = cant_total // lotes + (1 if k < cant_total % lotes else 0)
            f_hab = laborable(ini + timedelta(days=7 + int(14 * n / 45) + random.randint(0, 2)))
            f_arm = laborable(f_hab + timedelta(days=random.randint(2, 4)))
            f_sol = laborable(f_arm + timedelta(days=random.randint(2, 3)))
            f_lib = laborable(f_sol + timedelta(days=random.randint(1, 2)))
            pasada = lambda d: d.isoformat() if d <= hoy else None
            piezas.append({'Código OT': OT, 'N° item': n, 'Tipo de elemento': tipo, 'Conjunto / marca': f'{tipo[:3]}-{k + 1:02d}',
                           'Perfil': perfil, 'Longitud (mm)': random.choice([6000, 7500, 9000, 9850, 12000]), 'Cantidad': cant,
                           'Peso unit. neto (kg)': pu, 'Área total (m²)': round(pu * cant * 0.014, 1), 'Contratista': contr[n % 4],
                           'O.F.': 5 + n // 8, 'P.U. (S/ por kg)': 0.6 if pu > 250 else 1.1,
                           'F. habilitado': pasada(f_hab), 'F. armado': pasada(f_arm) if f_hab <= hoy else None,
                           'F. soldeo': pasada(f_sol) if f_arm <= hoy else None, 'F. liberación calidad': pasada(f_lib) if f_sol <= hoy else None})
    wb['PROYECTO'].cell(3, 7, sum(p['Cantidad'] for p in piezas))          # N° piezas
    escribir(wb['PIEZAS'], piezas)

    procesos = ['01 Ingeniería', '02 Compra de material', '03 Habilitado - sierra cinta', '04 Habilitado - cizalla-punzonadora',
                '05 Habilitado - mesa CNC', '06 Roscado de barra lisa', '07 Doblez (servicio externo)', '08 Armado', '09 Soldeo',
                '10 Limpieza', '11 Despacho a pintura', '12 Granallado y pintura (externo)', '13 Despacho a obra']
    hh_est = [85, 9, 33, 25, 27, 17, 10, 340, 420, 170, 47, 106, 28]
    escribir(wb['COTIZ_PROCESO'], [{'Código OT': OT, 'Proceso': p, 'HH estimadas': h, 'Estimado por': 'Cotizador (demo)',
                                    'Fecha de estimación': (ini - timedelta(days=12)).isoformat()} for p, h in zip(procesos, hh_est)])
    escribir(wb['PROCESO_FECHAS'], [{'Código OT': OT, 'Proceso': '01 Ingeniería', 'Fecha inicio real': ini.isoformat(),
                                     'Fecha fin real': laborable(ini + timedelta(days=9)).isoformat()},
                                    {'Código OT': OT, 'Proceso': '08 Armado', 'Fecha inicio real': laborable(ini + timedelta(days=10)).isoformat()}])
    escribir(wb['TAREO'], [{'Fecha': laborable(hoy - timedelta(days=2)).isoformat() if laborable(hoy - timedelta(days=2)) <= hoy else hoy.isoformat(),
                            'Código OT': OT, 'Proceso': '09 Soldeo', 'Contratista / cuadrilla': 'T&C FABRICACION', 'N° personas': 6,
                            'Horas por persona': 8, 'Kg avanzados': 2100, 'Observación (falta de material, retrabajo, etc.)': 'DEMO'}])
    escribir(wb['COMPRAS'], [{'Código OT': OT, 'N° OC': f'OC-DEMO-{k}', 'Proveedor': 'Proveedor acero (demo)', 'Perfil / material': m, 'Kg': kg,
                              'Fecha pedido': (ini + timedelta(days=d)).isoformat(), 'Fecha prometida': (ini + timedelta(days=d + 5)).isoformat(),
                              'Fecha recepción': (ini + timedelta(days=d + 5 + a)).isoformat()}
                             for k, (m, kg, d, a) in enumerate([('W8x31', 10200, 2, 0), ('W10x22', 8600, 2, 2), ('Ángulos L', 14800, 3, 4)], 1)])
    escribir(wb['SERV_EXTERNOS'], [{'Código OT': OT, 'Servicio': 'Granallado y pintura', 'Proveedor': 'Pintura industrial (demo)',
                                    'Fecha envío': laborable(hoy - timedelta(days=1)).isoformat(), 'Kg': 9800, 'm²': 210}])
    escribir(wb['EVENTOS'], [{'Código OT': OT, 'Fecha': (ini + timedelta(days=11)).isoformat(), 'Tipo de evento': 'Retraso de información del cliente (RFI)',
                              'Imputable a': 'Cliente', 'Días de impacto': 2, '¿Se pidió ampliación?': 'Sí',
                              'Descripción': 'DEMO: el cliente demoró en confirmar el espesor de placas base'}])
    escribir(wb['PARADAS'], [{'Máquina': 'Cizalladora Punzonadora Pedimax', 'Fecha y hora inicio': f'{laborable(ini + timedelta(days=12))} 09:00',
                              'Fecha y hora fin': f'{laborable(ini + timedelta(days=12))} 11:30', '¿Planificada?': 'No', 'Causa': 'Cambio de herramienta / setup',
                              'Código OT afectada': OT, 'Observación': 'DEMO'}])
    wb['LEEME']['A1'] = 'EJEMPLO LLENO (DATOS DEMO) · formato único v1 · proyecto OT-DEMO-001'
    wb.save(SALIDA)
    print('OK', SALIDA, '| piezas:', len(piezas), 'filas,', sum(p['Cantidad'] * p['Peso unit. neto (kg)'] for p in piezas) / 1000, 't')


if __name__ == '__main__':
    main()
