"""Genera entregables/Formato_Unico_Steelser_v1.xlsx: el formato único de registro (Componente 1).

Sigue las convenciones de los Excel de la empresa (control por pieza OPE-PRO-FR, clases de peso
Liviana/Mediana/Pesada, tipos de elemento, contratistas por O.F.) y agrega lo que el modelo necesita:
tareo diario, fechas por proceso, paradas, servicios externos, compras, eventos y la estimación del cotizador.
Uso: py scripts/generar_formato_unico.py
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
SALIDA = RAIZ / 'entregables' / 'Formato_Unico_Steelser_v1.xlsx'

AZUL, CELESTE, CEL_CL, ACERO, NARANJA = '0A6ED1', '1B90FF', 'E8F3FF', '223548', 'E9730C'
fino = Side(style='thin', color='C9CED3')
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)

from plataforma.formato import HOJAS, LISTAS  # noqa: E402


def estilo_cab(c, color=AZUL):
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=color)
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    c.border = BORDE


def main(salida=SALIDA):
    wb = Workbook()
    lee = wb.active
    lee.title = 'LEEME'
    lee['A1'] = 'FORMATO ÚNICO DE REGISTRO · STEELSER S.A.C. · v1'
    lee['A1'].font = Font(bold=True, size=16, color=ACERO)
    lee['A2'] = 'Componente 1 de la tesis: datos estandarizados para estimar horas, disponibilidad y plazos. Basado en el control por pieza OPE-PRO-FR que ya usa la planta.'
    lee['A2'].alignment = Alignment(wrap_text=True)
    lee.merge_cells('A2:E2')
    lee.row_dimensions[2].height = 32
    filas = [('Hoja', 'Quién la llena · cuándo', 'Para qué')] + [(h[0], h[1], h[2]) for h in HOJAS]
    for i, f in enumerate(filas, start=4):
        for j, v in enumerate(f, start=1):
            c = lee.cell(i, j, v)
            if i == 4:
                estilo_cab(c, ACERO)
            else:
                c.border = BORDE
                c.alignment = Alignment(wrap_text=True, vertical='top')
    lee.column_dimensions['A'].width = 18
    lee.column_dimensions['B'].width = 42
    lee.column_dimensions['C'].width = 70
    r = len(filas) + 6
    reglas = ['REGLAS', '1. Una fila por registro; no combinar celdas ni dejar filas en blanco intermedias.',
              '2. Usar las listas desplegables (celdas celestes) en vez de escribir a mano.',
              '3. Fechas en formato AAAA-MM-DD; fecha y hora como AAAA-MM-DD HH:MM.',
              '4. La fila 3 de cada hoja (gris) es un EJEMPLO: borrarla antes de cargar los datos reales.',
              '5. El Código OT debe ser el mismo en todas las hojas.',
              '6. Las columnas naranjas se calculan solas (HH, peso total, clase de peso, horas, días).',
              '7. Tareo: registrar todos los días, también el personal del contratista (no basta la valorización por kg).',
              '8. Paradas: registrar cualquier parada de más de 15 minutos de las 4 máquinas, planificada o no.',
              '9. Confidencial: uso exclusivo para la tesis según la autorización firmada por la empresa.']
    for k, t in enumerate(reglas):
        c = lee.cell(r + k, 1, t)
        c.font = Font(bold=(k == 0), color=ACERO)
        lee.merge_cells(start_row=r + k, start_column=1, end_row=r + k, end_column=3)

    ls = wb.create_sheet('LISTAS')
    rangos = {}
    for j, (nombre, valores) in enumerate(LISTAS.items(), start=1):
        estilo_cab(ls.cell(1, j, nombre), ACERO)
        for i, v in enumerate(valores, start=2):
            ls.cell(i, j, v)
        col = get_column_letter(j)
        rangos[nombre] = f'LISTAS!${col}$2:${col}${len(valores) + 1}'
        ls.column_dimensions[col].width = max(14, max(len(v) for v in valores) + 2)

    for nombre, quien, para, cols, ejemplo in HOJAS:
        ws = wb.create_sheet(nombre)
        ws['A1'] = f'{nombre} · {para}  ({quien})'
        ws['A1'].font = Font(bold=True, color=ACERO, size=12)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(cols), 12))
        for j, (titulo, ancho, tipo) in enumerate(cols, start=1):
            c = ws.cell(2, j, titulo)
            estilo_cab(c, NARANJA if tipo and tipo.startswith('formula') else AZUL)
            ws.column_dimensions[get_column_letter(j)].width = ancho
        ws.row_dimensions[2].height = 42
        ws.freeze_panes = 'A3'
        n_filas = 500
        cols_idx = {t: get_column_letter(j) for j, (t, _, _) in enumerate(cols, start=1)}
        for j, (titulo, ancho, tipo) in enumerate(cols, start=1):
            L = get_column_letter(j)
            rango = f'{L}3:{L}{n_filas}'
            if tipo in LISTAS:
                dv = DataValidation(type='list', formula1=f'={rangos[tipo]}', allow_blank=True, showErrorMessage=True,
                                    errorTitle='Valor no válido', error='Elija un valor de la lista desplegable.')
                ws.add_data_validation(dv)
                dv.add(rango)
                for i in range(3, 60):
                    ws.cell(i, j).fill = PatternFill('solid', fgColor=CEL_CL)
            elif tipo == 'int':
                dv = DataValidation(type='whole', operator='between', formula1='0', formula2='100000', allow_blank=True,
                                    showErrorMessage=True, error='Número entero.')
                ws.add_data_validation(dv); dv.add(rango)
            elif tipo == 'num':
                dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                                    showErrorMessage=True, error='Número mayor o igual a 0.')
                ws.add_data_validation(dv); dv.add(rango)
            elif tipo == 'horas':
                dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='12', allow_blank=True,
                                    showErrorMessage=True, error='Entre 0 y 12 horas por persona.')
                ws.add_data_validation(dv); dv.add(rango)
            elif tipo == 'pct':
                dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='1', allow_blank=True,
                                    showErrorMessage=True, error='Fracción entre 0 y 1 (17 % = 0.17).')
                ws.add_data_validation(dv); dv.add(rango)
                for i in range(3, n_filas):
                    ws.cell(i, j).number_format = '0%'
            elif tipo and tipo.startswith('formula'):
                for i in range(3, n_filas):
                    f = None
                    if tipo == 'formula_hh':
                        f = f'=IF(AND({cols_idx["N° personas"]}{i}<>"",{cols_idx["Horas por persona"]}{i}<>""),{cols_idx["N° personas"]}{i}*{cols_idx["Horas por persona"]}{i},"")'
                    elif tipo == 'formula_peso':
                        f = f'=IF(AND({cols_idx["Cantidad"]}{i}<>"",{cols_idx["Peso unit. neto (kg)"]}{i}<>""),{cols_idx["Cantidad"]}{i}*{cols_idx["Peso unit. neto (kg)"]}{i},"")'
                    elif tipo == 'formula_clase':
                        p = f'{cols_idx["Peso unit. neto (kg)"]}{i}'
                        f = f'=IF({p}="","",IF({p}<=250,"Liviana (0-250 kg)",IF({p}<=1000,"Mediana (251-1000 kg)","Pesada (>1000 kg)")))'
                    elif tipo == 'formula_horas':
                        a, b = cols_idx['Fecha y hora inicio'], cols_idx['Fecha y hora fin']
                        fh = lambda x: f'IF(ISNUMBER({x}{i}),{x}{i},DATEVALUE(LEFT({x}{i},10))+TIMEVALUE(MID({x}{i},12,5)))'
                        f = f'=IF(AND({a}{i}<>"",{b}{i}<>""),ROUND(({fh(b)}-{fh(a)})*24,1),"")'
                    elif tipo in ('formula_dias', 'formula_atraso'):
                        a, b = (cols_idx['Fecha envío'], cols_idx['Fecha retorno']) if tipo == 'formula_dias' else \
                               (cols_idx['Fecha prometida'], cols_idx['Fecha recepción'])
                        fd = lambda x: f'IF(ISNUMBER({x}{i}),{x}{i},DATEVALUE({x}{i}))'
                        dif = f'{fd(b)}-{fd(a)}'
                        f = f'=IF(AND({a}{i}<>"",{b}{i}<>""),{dif if tipo == "formula_dias" else f"MAX(0,{dif})"},"")'
                    c = ws.cell(i, j, f)
                    c.fill = PatternFill('solid', fgColor='FFF3E6')
        for j, v in enumerate(ejemplo, start=1):
            c = ws.cell(3, j)
            if v is not None and not (isinstance(c.value, str) and str(c.value).startswith('=')):
                c.value = v
            c.font = Font(italic=True, color='6A6D70')
            if not (cols[j - 1][2] or '').startswith('formula'):
                c.fill = PatternFill('solid', fgColor='EEF1F3')
        ws.cell(3, len(cols) + 1, '← EJEMPLO: borrar esta fila').font = Font(italic=True, color=NARANJA, bold=True)

    wb.move_sheet('LISTAS', offset=len(wb.sheetnames))
    salida = Path(salida)
    salida.parent.mkdir(exist_ok=True)
    wb.save(salida)
    print('OK', salida)
    return salida


if __name__ == '__main__':
    main()
