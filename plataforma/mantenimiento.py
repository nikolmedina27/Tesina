"""SteelPlan · Mantenimiento de máquinas: fichas, indicadores TPM, paradas, plan preventivo, órdenes y Gantt.

Lee y escribe en data/plataforma.db (tablas `maquina`, `plan_mantenimiento`, `orden_mantenimiento` y `parada`).
Los cálculos están en `dss/mantenimiento.py`. Las rutas se registran desde server.py con `registrar(...)`.

Reglas:
- Al cerrar una orden con inicio y fin se crea su parada (planificada si es preventiva), salvo que ya venga de una.
- Las órdenes preventivas se generan solas para los planes que vencen en los próximos días (idempotente: un plan
  nunca tiene más de una orden abierta).
"""
from datetime import date, datetime, timedelta
from io import BytesIO

import numpy as np
from fastapi import Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from dss.crp_engine import WEEKMASK
from dss.mantenimiento import META_DK, dk_mensual, indicadores, proximo_vencimiento, semaforo

HORIZONTE_DIAS = 10          # días laborables hacia adelante en que se crean las órdenes preventivas
EDITAN = ('jefe_taller', 'mantenimiento')
REPORTAN = ('jefe_taller', 'mantenimiento', 'supervisor', 'calidad')
TIPOS_ORDEN = ('PREVENTIVO', 'CORRECTIVO', 'PREDICTIVO')
ESTADOS_ORDEN = ('PENDIENTE', 'EN_CURSO', 'CERRADA', 'ANULADA')


class FichaIn(BaseModel):
    centro: str | None = None
    marca: str | None = None
    modelo: str | None = None
    anio: int | None = Field(None, ge=1950, le=2100)
    criticidad: str | None = Field(None, pattern='^[ABC]$')
    horas_turno: float | None = Field(None, gt=0, le=24)
    turnos: int | None = Field(None, ge=1, le=3)
    descripcion: str | None = None


class PlanIn(BaseModel):
    maquina_id: int
    tarea: str = Field(min_length=3, max_length=200)
    frecuencia_dias: int | None = Field(None, ge=1, le=730)
    frecuencia_horas: float | None = Field(None, gt=0, le=20000)
    duracion_h: float = Field(1, gt=0, le=200)
    responsable: str | None = None


class PlanCambio(BaseModel):
    tarea: str | None = None
    frecuencia_dias: int | None = Field(None, ge=1, le=730)
    frecuencia_horas: float | None = Field(None, gt=0, le=20000)
    duracion_h: float | None = Field(None, gt=0, le=200)
    responsable: str | None = None
    activo: bool | None = None


class OrdenIn(BaseModel):
    maquina_id: int
    tipo: str = 'CORRECTIVO'
    titulo: str = Field(min_length=3, max_length=200)
    descripcion: str | None = None
    programada_para: str | None = None
    duracion_h: float | None = Field(None, gt=0, le=200)
    tecnico: str | None = None
    parada_id: int | None = None
    iniciar: bool = False           # la máquina quedó fuera de servicio ahora: la orden arranca EN_CURSO


class OrdenCambio(BaseModel):
    estado: str | None = None
    programada_para: str | None = None
    duracion_h: float | None = Field(None, gt=0, le=200)
    inicio: str | None = None
    fin: str | None = None
    tecnico: str | None = None
    repuestos: str | None = None
    costo: float | None = Field(None, ge=0)
    descripcion: str | None = None
    causa: str | None = None        # causa de la parada que se crea al cerrar una correctiva


def _ahora():
    return datetime.now().strftime('%Y-%m-%d %H:%M')


def _fecha(s, nombre):
    try:
        return datetime.fromisoformat(str(s).replace('T', ' ')).strftime('%Y-%m-%d %H:%M' if len(str(s)) > 10 else '%Y-%m-%d')
    except ValueError:
        raise HTTPException(422, f'{nombre} no válida')


def registrar(app, plat, usuario_actual, requiere, actividad, causas, ops_maquinas):
    """ops_maquinas(con) → operaciones planificadas de los proyectos en curso por máquina (para el Gantt)."""

    def maquinas(p, solo_activas=True):
        q = 'SELECT * FROM maquina' + (' WHERE activa=1' if solo_activas else '') + ' ORDER BY id'
        return [dict(r) for r in p.execute(q)]

    def maquina(p, mid):
        r = p.execute('SELECT * FROM maquina WHERE id=?', (mid,)).fetchone()
        if not r:
            raise HTTPException(404, 'Máquina no encontrada')
        return dict(r)

    def paradas_de(p, nombre, desde=None):
        q, a = 'SELECT * FROM parada WHERE maquina=?', [nombre]
        if desde:
            q += ' AND fin>=?'
            a.append(desde)
        return [dict(r) for r in p.execute(q + ' ORDER BY inicio', a)]

    def inicio_registro(p):
        """Fecha desde la que se registran paradas (la primera registrada en la planta); antes no hay dato."""
        r = p.execute('SELECT MIN(inicio) FROM parada').fetchone()[0]
        return r[:10] if r else None

    def desde(p, dias):
        d = str(date.today() - timedelta(days=dias))
        r0 = inicio_registro(p)
        return max(d, r0) if r0 else d

    def sin_dato(p, i):
        """Si la planta todavía no registró ninguna parada, no hay disponibilidad que informar (no es 100 %)."""
        return i if inicio_registro(p) else i | dict(dk=None, mtbf_h=None, mttr_h=None)

    def ultima_ejecucion(p, plan_id):
        r = p.execute("SELECT MAX(COALESCE(fin, cerrado_en)) FROM orden_mantenimiento WHERE plan_id=? AND estado='CERRADA'", (plan_id,)).fetchone()
        return r[0][:10] if r and r[0] else None

    def generar_ordenes(p, u=None):
        """Crea la orden preventiva PENDIENTE de cada plan que vence dentro del horizonte (o ya venció)."""
        hoy = str(date.today())
        creadas = 0
        for pl in p.execute('''SELECT pm.*, m.nombre maquina FROM plan_mantenimiento pm JOIN maquina m ON m.id=pm.maquina_id
                               WHERE pm.activo=1 AND m.activa=1''').fetchall():
            pl = dict(pl)
            if p.execute("SELECT 1 FROM orden_mantenimiento WHERE plan_id=? AND estado IN ('PENDIENTE','EN_CURSO')", (pl['id'],)).fetchone():
                continue
            vence, falta = proximo_vencimiento(pl, ultima_ejecucion(p, pl['id']), hoy)
            if vence is None or falta > HORIZONTE_DIAS:
                continue
            p.execute('''INSERT INTO orden_mantenimiento (maquina_id, plan_id, tipo, titulo, programada_para, duracion_h, tecnico, creado_por)
                         VALUES (?,?, 'PREVENTIVO', ?,?,?,?,?)''',
                      (pl['maquina_id'], pl['id'], pl['tarea'], vence, pl['duracion_h'], pl['responsable'], u['id'] if u else None))
            creadas += 1
        if creadas:
            p.commit()
        return creadas

    def resumen_maquina(p, m, hoy, ahora):
        par = paradas_de(p, m['nombre'], str(date.today() - timedelta(days=400)))
        i30 = sin_dato(p, indicadores(par, desde(p, 30), str(date.today() + timedelta(days=1)), m['horas_turno'] or 8, m['turnos'] or 1))
        abierta = next((x for x in par if x['inicio'] <= ahora < x['fin']), None)
        en_curso = p.execute("SELECT id, titulo, tipo FROM orden_mantenimiento WHERE maquina_id=? AND estado='EN_CURSO' ORDER BY tipo='CORRECTIVO' DESC",
                             (m['id'],)).fetchone()
        ult = p.execute('''SELECT MAX(COALESCE(fin, cerrado_en)) FROM orden_mantenimiento WHERE maquina_id=? AND tipo='PREVENTIVO'
                           AND estado='CERRADA' ''', (m['id'],)).fetchone()[0]
        prox = p.execute('''SELECT id, titulo, programada_para FROM orden_mantenimiento WHERE maquina_id=? AND tipo='PREVENTIVO'
                            AND estado IN ('PENDIENTE','EN_CURSO') ORDER BY programada_para LIMIT 1''', (m['id'],)).fetchone()
        vencidas = p.execute('''SELECT COUNT(*) FROM orden_mantenimiento WHERE maquina_id=? AND estado='PENDIENTE'
                                AND programada_para < ?''', (m['id'], hoy)).fetchone()[0]
        falla = (abierta and not abierta['planificada']) or (en_curso and en_curso['tipo'] == 'CORRECTIVO')
        estado = 'PARADA' if falla else 'MANTENIMIENTO' if (abierta or en_curso) else 'OPERANDO'
        return m | dict(estado=estado, parada_actual=abierta, orden_en_curso=dict(en_curso) if en_curso else None,
                        ultimo_preventivo=ult[:10] if ult else None, proximo_preventivo=dict(prox) if prox else None,
                        preventivos_vencidos=vencidas, indicadores_30d=i30,
                        semaforo=semaforo(i30['dk'], vencidas, estado == 'PARADA'))

    # ---------------------------------------------------------------- máquinas
    def resumen_todas(p):
        hoy, ahora = str(date.today()), _ahora()
        return [resumen_maquina(p, m, hoy, ahora) for m in maquinas(p)]

    @app.get('/api/mant/maquinas')
    def lista_maquinas(u=Depends(usuario_actual), p=Depends(plat)):
        generar_ordenes(p)
        out = resumen_todas(p)
        dks = [m['indicadores_30d']['dk'] for m in out if m['indicadores_30d']['dk'] is not None]
        return dict(maquinas=out, meta_dk=META_DK, resumen=dict(
            dk_promedio=round(float(np.mean(dks)), 4) if dks else None,
            fallas_30d=sum(m['indicadores_30d']['n_fallas'] for m in out),
            horas_perdidas_30d=round(sum(m['indicadores_30d']['horas_parada'] for m in out), 1),
            preventivos_vencidos=sum(m['preventivos_vencidos'] for m in out),
            paradas_ahora=sum(m['estado'] != 'OPERANDO' for m in out)))

    @app.get('/api/mant/maquinas/{mid}')
    def ficha(mid: int, u=Depends(usuario_actual), p=Depends(plat)):
        generar_ordenes(p)
        m = maquina(p, mid)
        hoy = date.today()
        par = paradas_de(p, m['nombre'])
        ht, tu = m['horas_turno'] or 8, m['turnos'] or 1
        manana = str(hoy + timedelta(days=1))
        ventanas = {k: sin_dato(p, indicadores(par, desde(p, d), manana, ht, tu)) | dict(desde=desde(p, d)) for k, d in (('30d', 30), ('90d', 90), ('365d', 365))}
        planes = []
        for pl in p.execute('SELECT * FROM plan_mantenimiento WHERE maquina_id=? ORDER BY activo DESC, id', (mid,)):
            pl = dict(pl)
            ult = ultima_ejecucion(p, pl['id'])
            vence, falta = proximo_vencimiento(pl, ult, str(hoy))
            planes.append(pl | dict(ultima=ult, vence=vence, dias_para_vencer=falta))
        ordenes = [dict(r) for r in p.execute('SELECT * FROM orden_mantenimiento WHERE maquina_id=? ORDER BY COALESCE(fin, programada_para, creado_en) DESC', (mid,))]
        return dict(maquina=resumen_maquina(p, m, str(hoy), _ahora()), indicadores=ventanas, dk_mensual=dk_mensual(par, str(hoy), 12, ht, tu, inicio_registro(p) or str(hoy + timedelta(days=1))), inicio_registro=inicio_registro(p),
                    paradas=[x | dict(horas=round((datetime.fromisoformat(x['fin']) - datetime.fromisoformat(x['inicio'])).total_seconds() / 3600, 1))
                             for x in reversed(par)], ordenes=ordenes, planes=planes, meta_dk=META_DK)

    @app.patch('/api/mant/maquinas/{mid}')
    def editar_ficha(mid: int, d: FichaIn, u=Depends(requiere(*EDITAN)), p=Depends(plat)):
        maquina(p, mid)
        cambios = {k: v for k, v in d.model_dump().items() if v is not None}
        if cambios:
            p.execute(f'UPDATE maquina SET {", ".join(k + "=?" for k in cambios)} WHERE id=?', (*cambios.values(), mid))
            p.commit()
        return {'ok': True}

    # ---------------------------------------------------------------- paradas (todas)
    def filtrar_paradas(p, desde, hasta, maquina_, causa, tipo, proyecto_id):
        q = '''SELECT pa.*, us.nombre usuario, pr.codigo proyecto, ROUND((julianday(pa.fin)-julianday(pa.inicio))*24, 2) horas
               FROM parada pa LEFT JOIN usuario us ON us.id=pa.registrado_por LEFT JOIN proyecto_activo pr ON pr.id=pa.proyecto_id WHERE 1=1'''
        a = []
        if desde:
            q += ' AND pa.fin >= ?'
            a.append(desde)
        if hasta:
            q += ' AND pa.inicio < ?'
            a.append(str(date.fromisoformat(hasta[:10]) + timedelta(days=1)))
        for col, v in (('pa.maquina', maquina_), ('pa.causa', causa), ('pa.proyecto_id', proyecto_id)):
            if v:
                q += f' AND {col}=?'
                a.append(v)
        if tipo in ('falla', 'planificada'):
            q += ' AND pa.planificada=?'
            a.append(int(tipo == 'planificada'))
        return [dict(r) for r in p.execute(q + ' ORDER BY pa.inicio DESC', a)]

    @app.get('/api/mant/paradas')
    def paradas(desde: str = '', hasta: str = '', maquina: str = '', causa: str = '', tipo: str = '', proyecto_id: int | None = None,
                u=Depends(usuario_actual), p=Depends(plat)):
        filas = filtrar_paradas(p, desde, hasta, maquina, causa, tipo, proyecto_id)
        por = lambda k: sorted(({'clave': c, 'horas': round(sum(f['horas'] or 0 for f in filas if (f[k] or '—') == c), 1),
                                  'n': sum(1 for f in filas if (f[k] or '—') == c)} for c in {f[k] or '—' for f in filas}), key=lambda x: -x['horas'])
        return dict(filas=filas, total_horas=round(sum(f['horas'] or 0 for f in filas), 1), n=len(filas),
                    n_fallas=sum(1 for f in filas if not f['planificada']), por_maquina=por('maquina'), por_causa=por('causa'))

    @app.get('/api/mant/paradas/excel')
    def paradas_excel(desde: str = '', hasta: str = '', maquina: str = '', causa: str = '', tipo: str = '', proyecto_id: int | None = None,
                      u=Depends(usuario_actual), p=Depends(plat)):
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        filas = filtrar_paradas(p, desde, hasta, maquina, causa, tipo, proyecto_id)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'PARADAS'
        cols = [('Máquina', 'maquina', 30), ('Inicio', 'inicio', 17), ('Fin', 'fin', 17), ('Horas', 'horas', 8), ('Tipo', None, 12),
                ('Causa', 'causa', 28), ('OT', 'proyecto', 14), ('Orden mant.', 'orden_id', 11), ('Observación', 'observacion', 40), ('Registró', 'usuario', 22)]
        ws.append([c[0] for c in cols])
        for j, (_, _, w) in enumerate(cols, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
            ws.cell(1, j).font = Font(bold=True, color='FFFFFF')
            ws.cell(1, j).fill = PatternFill('solid', fgColor='0B4B8F')
        for f in filas:
            ws.append([('Planificada' if f['planificada'] else 'Falla') if k is None else f[k] for _, k, _ in cols])
        ws.freeze_panes = 'A2'
        b = BytesIO()
        wb.save(b)
        b.seek(0)
        nombre = f'paradas_{desde or "inicio"}_{hasta or date.today()}.xlsx'
        return StreamingResponse(b, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                 headers={'Content-Disposition': f'attachment; filename="{nombre}"'})

    # ---------------------------------------------------------------- plan preventivo
    @app.get('/api/mant/planes')
    def planes(u=Depends(usuario_actual), p=Depends(plat)):
        hoy = str(date.today())
        out = []
        for pl in p.execute('SELECT pm.*, m.nombre maquina FROM plan_mantenimiento pm JOIN maquina m ON m.id=pm.maquina_id ORDER BY m.id, pm.id'):
            pl = dict(pl)
            ult = ultima_ejecucion(p, pl['id'])
            vence, falta = proximo_vencimiento(pl, ult, hoy)
            out.append(pl | dict(ultima=ult, vence=vence, dias_para_vencer=falta))
        return out

    @app.post('/api/mant/planes')
    def crear_plan(d: PlanIn, u=Depends(requiere(*EDITAN)), p=Depends(plat)):
        m = maquina(p, d.maquina_id)
        if not d.frecuencia_dias and not d.frecuencia_horas:
            raise HTTPException(422, 'Indique la frecuencia en días o en horas de uso')
        cur = p.execute('''INSERT INTO plan_mantenimiento (maquina_id, tarea, frecuencia_dias, frecuencia_horas, duracion_h, responsable, desde)
                           VALUES (?,?,?,?,?,?,?)''', (d.maquina_id, d.tarea, d.frecuencia_dias, d.frecuencia_horas, d.duracion_h, d.responsable, str(date.today())))
        actividad(p, u, 'mantenimiento', f'Agregó al plan preventivo de {m["nombre"]}: {d.tarea}')
        p.commit()
        generar_ordenes(p, u)
        return {'id': cur.lastrowid}

    @app.patch('/api/mant/planes/{plid}')
    def cambiar_plan(plid: int, d: PlanCambio, u=Depends(requiere(*EDITAN)), p=Depends(plat)):
        if not p.execute('SELECT 1 FROM plan_mantenimiento WHERE id=?', (plid,)).fetchone():
            raise HTTPException(404)
        cambios = {k: (int(v) if isinstance(v, bool) else v) for k, v in d.model_dump().items() if v is not None}
        if cambios:
            p.execute(f'UPDATE plan_mantenimiento SET {", ".join(k + "=?" for k in cambios)} WHERE id=?', (*cambios.values(), plid))
            if cambios.get('activo') == 0:      # un plan desactivado no deja órdenes pendientes colgando
                p.execute("UPDATE orden_mantenimiento SET estado='ANULADA', cerrado_en=? WHERE plan_id=? AND estado='PENDIENTE'", (_ahora(), plid))
            p.commit()
        return {'ok': True}

    # ---------------------------------------------------------------- órdenes de mantenimiento
    @app.get('/api/mant/ordenes')
    def ordenes(estado: str = '', maquina_id: int | None = None, u=Depends(usuario_actual), p=Depends(plat)):
        generar_ordenes(p)
        q = '''SELECT o.*, m.nombre maquina, us.nombre creado_por_nombre FROM orden_mantenimiento o JOIN maquina m ON m.id=o.maquina_id
               LEFT JOIN usuario us ON us.id=o.creado_por WHERE 1=1'''
        a = []
        if estado:
            q += ' AND o.estado=?'
            a.append(estado)
        if maquina_id:
            q += ' AND o.maquina_id=?'
            a.append(maquina_id)
        hoy = str(date.today())
        filas = [dict(r) for r in p.execute(q + " ORDER BY o.estado='CERRADA', o.estado='ANULADA', COALESCE(o.programada_para, o.creado_en)", a)]
        for f in filas:
            f['vencida'] = f['estado'] == 'PENDIENTE' and bool(f['programada_para']) and f['programada_para'][:10] < hoy
        return filas

    @app.post('/api/mant/ordenes')
    def crear_orden(d: OrdenIn, u=Depends(requiere(*REPORTAN)), p=Depends(plat)):
        m = maquina(p, d.maquina_id)
        if d.tipo not in TIPOS_ORDEN:
            raise HTTPException(422, 'Tipo de orden no válido')
        if d.parada_id:
            pa = p.execute('SELECT * FROM parada WHERE id=?', (d.parada_id,)).fetchone()
            if not pa or pa['maquina'] != m['nombre']:
                raise HTTPException(422, 'La parada no existe o es de otra máquina')
        cur = p.execute('''INSERT INTO orden_mantenimiento (maquina_id, tipo, titulo, descripcion, programada_para, duracion_h, tecnico, parada_id,
                           creado_por, estado, inicio) VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
                        (d.maquina_id, d.tipo, d.titulo, d.descripcion, _fecha(d.programada_para, 'Fecha') if d.programada_para else None,
                         d.duracion_h, d.tecnico, d.parada_id, u['id'], 'EN_CURSO' if d.iniciar else 'PENDIENTE', _ahora() if d.iniciar else None))
        if d.parada_id:
            p.execute('UPDATE parada SET orden_id=? WHERE id=?', (cur.lastrowid, d.parada_id))
        actividad(p, u, 'mantenimiento', f'Abrió orden {d.tipo.lower()} OM-{cur.lastrowid} en {m["nombre"]}: {d.titulo}')
        p.commit()
        return {'id': cur.lastrowid}

    @app.patch('/api/mant/ordenes/{oid}')
    def cambiar_orden(oid: int, d: OrdenCambio, u=Depends(requiere(*EDITAN)), p=Depends(plat)):
        o = p.execute('SELECT o.*, m.nombre maquina FROM orden_mantenimiento o JOIN maquina m ON m.id=o.maquina_id WHERE o.id=?', (oid,)).fetchone()
        if not o:
            raise HTTPException(404)
        o = dict(o)
        if o['estado'] in ('CERRADA', 'ANULADA') and d.estado not in (None, o['estado']):
            raise HTTPException(422, 'La orden ya está cerrada')
        cambios = {k: v for k, v in d.model_dump().items() if v is not None and k != 'causa'}
        if 'estado' in cambios and cambios['estado'] not in ESTADOS_ORDEN:
            raise HTTPException(422, 'Estado no válido')
        for k in ('inicio', 'fin', 'programada_para'):
            if k in cambios:
                cambios[k] = _fecha(cambios[k], k.replace('_', ' ').capitalize())
        if cambios.get('estado') == 'EN_CURSO' and not (cambios.get('inicio') or o['inicio']):
            cambios['inicio'] = _ahora()
        if cambios.get('estado') == 'CERRADA':
            ini, fin = cambios.get('inicio') or o['inicio'], cambios.get('fin') or o['fin'] or _ahora()
            if not ini:
                raise HTTPException(422, 'Para cerrar la orden indique cuándo empezó')
            if fin <= ini:
                raise HTTPException(422, 'El fin debe ser posterior al inicio')
            if fin > _ahora():
                raise HTTPException(422, 'El fin no puede ser futuro')
            cambios.update(inicio=ini, fin=fin, cerrado_en=_ahora())
            if not o['parada_id']:          # la intervención dejó la máquina fuera de servicio: queda como parada
                planificada = o['tipo'] != 'CORRECTIVO'
                causa = 'Mantenimiento preventivo' if planificada else (d.causa if d.causa in causas else 'Falla mecánica')
                pid = p.execute('''INSERT INTO parada (maquina, inicio, fin, planificada, causa, observacion, registrado_por, orden_id)
                                   VALUES (?,?,?,?,?,?,?,?)''', (o['maquina'], ini, fin, int(planificada), causa, f'OM-{oid}: {o["titulo"]}', u['id'], oid)).lastrowid
                cambios['parada_id'] = pid
        if cambios.get('estado') == 'ANULADA':
            cambios['cerrado_en'] = _ahora()
        if cambios:
            p.execute(f'UPDATE orden_mantenimiento SET {", ".join(k + "=?" for k in cambios)} WHERE id=?', (*cambios.values(), oid))
        if 'estado' in cambios:
            actividad(p, u, 'mantenimiento', f'OM-{oid} ({o["maquina"]}) → {cambios["estado"].replace("_", " ").lower()}')
        p.commit()
        if cambios.get('estado') == 'CERRADA' and o['plan_id']:
            generar_ordenes(p, u)               # programa el siguiente preventivo de ese plan
        return {'ok': True}

    # ---------------------------------------------------------------- Gantt de máquinas
    @app.get('/api/mant/gantt')
    def gantt(desde: str = '', dias: int = 31, u=Depends(usuario_actual), p=Depends(plat)):
        generar_ordenes(p)
        ini = date.fromisoformat(desde[:10]) if desde else date.today() - timedelta(days=7)
        fin = ini + timedelta(days=max(1, min(dias, 400)))
        a, b = f'{ini} 00:00', f'{fin} 00:00'
        ms = maquinas(p)
        items = []
        for x in p.execute('SELECT pa.*, pr.codigo proyecto FROM parada pa LEFT JOIN proyecto_activo pr ON pr.id=pa.proyecto_id WHERE pa.fin>? AND pa.inicio<?', (a, b)):
            items.append(dict(maquina=x['maquina'], tipo='planificada' if x['planificada'] else 'falla', inicio=x['inicio'], fin=x['fin'],
                              titulo=x['causa'], detalle=x['observacion'] or '', proyecto=x['proyecto'], ref=f'parada:{x["id"]}'))
        ahora = _ahora()
        for o in p.execute('''SELECT o.*, m.nombre maquina FROM orden_mantenimiento o JOIN maquina m ON m.id=o.maquina_id
                              WHERE o.estado IN ('PENDIENTE','EN_CURSO') AND COALESCE(o.inicio, o.programada_para) IS NOT NULL'''):
            i0 = o['inicio'] or f'{o["programada_para"][:10]} 08:00'
            f0 = o['fin'] or (datetime.fromisoformat(i0) + timedelta(hours=o['duracion_h'] or 1)).strftime('%Y-%m-%d %H:%M')
            if o['estado'] == 'EN_CURSO' and f0 < ahora:
                f0 = ahora
            if f0 > a and i0 < b:
                items.append(dict(maquina=o['maquina'], tipo='orden_' + o['tipo'].lower(), inicio=i0, fin=f0, titulo=f'OM-{o["id"]} {o["titulo"]}',
                                  detalle=f'{o["estado"].replace("_", " ").lower()}{" · " + o["tecnico"] if o["tecnico"] else ""}', ref=f'orden:{o["id"]}',
                                  vencida=o['estado'] == 'PENDIENTE' and bool(o['programada_para']) and o['programada_para'][:10] < str(date.today())))
        for x in ops_maquinas(p):
            if x['fin'] > a and x['inicio'] < b:
                items.append(dict(x, tipo='proyecto'))
        return dict(desde=str(ini), hasta=str(fin), ahora=ahora, maquinas=[m['nombre'] for m in ms], items=items)

    return dict(resumen_maquinas=resumen_todas, generar_ordenes=generar_ordenes)
