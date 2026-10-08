"""Importa el formato único (Excel) a data/plataforma.db.

- Valida cada fila contra el formato (`plataforma.formato`): listas, fechas, rangos, orden de fechas, OT existente.
- Corre dentro de una transacción: en modo validar se revierte al final, así la validación es idéntica a la
  importación real (incluye OTs creadas en la hoja PROYECTO del mismo archivo).
- Es idempotente: reimportar el mismo archivo actualiza o ignora en vez de duplicar.
"""
import json
from datetime import date, datetime
from io import BytesIO

import openpyxl

from dss.datos import TIPOS
from .avance import ETAPAS, error_orden_etapas
from .formato import HOJAS, LISTAS

ORDEN_FECHAS_PIEZA = ['F. habilitado', 'F. armado', 'F. soldeo', 'F. liberación calidad', 'F. granallado', 'F. pintura', 'F. despacho']
COL_FECHA_PIEZA = dict(zip(ORDEN_FECHAS_PIEZA, [e[0] for e in ETAPAS]))
MAQUINAS = LISTAS['MAQUINAS']


class ErrorFila(ValueError):
    pass


# ------------------------------------------------------------------ conversión de celdas
def _fecha(v):
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        f = v
    else:
        s = str(v).strip()[:10]
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
            try:
                f = datetime.strptime(s, fmt).date()
                break
            except ValueError:
                f = None
        if f is None:
            raise ErrorFila(f'fecha no válida: "{v}" (use AAAA-MM-DD)')
    if not 2015 <= f.year <= 2035:
        raise ErrorFila(f'fecha fuera de rango: {f}')
    return f.isoformat()


def _fechahora(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M')
    s = str(v).strip().replace('T', ' ')
    for fmt in ('%Y-%m-%d %H:%M', '%Y-%m-%d %H:%M:%S', '%d/%m/%Y %H:%M'):
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d %H:%M')
        except ValueError:
            pass
    raise ErrorFila(f'fecha y hora no válida: "{v}" (use AAAA-MM-DD HH:MM)')


def _num(v, lo=0.0, hi=1e9, entero=False):
    try:
        x = float(str(v).replace(',', '.')) if not isinstance(v, (int, float)) else float(v)
    except ValueError:
        raise ErrorFila(f'número no válido: "{v}"')
    if not lo <= x <= hi:
        raise ErrorFila(f'valor {x:g} fuera de rango [{lo:g}, {hi:g}]')
    if entero:
        if x != int(x):
            raise ErrorFila(f'se esperaba un entero: {x:g}')
        return int(x)
    return x


def convertir(v, tipo):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if tipo is None:
        return str(v).strip()
    if tipo in LISTAS:
        s = str(v).strip().lower()
        for op in LISTAS[tipo]:
            if op.lower() == s:
                return op
        raise ErrorFila(f'"{v}" no está en la lista {tipo}')
    if tipo == 'fecha':
        return _fecha(v)
    if tipo == 'fechahora':
        return _fechahora(v)
    if tipo == 'int':
        return _num(v, 0, 1e6, entero=True)
    if tipo == 'num':
        return _num(v)
    if tipo == 'horas':
        return _num(v, 0, 12)
    if tipo == 'pct':
        x = _num(v, 0, 100)
        return x / 100 if x > 1 else x
    return None            # columnas calculadas (formula_*): se recalculan aquí


# ------------------------------------------------------------------ manejadores por hoja
def _ot(cx, codigo):
    if not codigo:
        raise ErrorFila('falta el Código OT')
    r = cx.execute('SELECT id FROM proyecto_activo WHERE codigo=?', (codigo.strip().upper(),)).fetchone()
    if not r:
        raise ErrorFila(f'la OT {codigo} no existe: créela en la hoja PROYECTO o desde el cotizador')
    return r[0]


def _requerir(f, *campos):
    faltan = [c for c in campos if f.get(c) is None]
    if faltan:
        raise ErrorFila('faltan datos: ' + ', '.join(faltan))


def h_proyecto(cx, f, u):
    _requerir(f, 'Código OT', 'Nombre del proyecto', 'Tipo de estructura', 'Toneladas metrado', 'Día 1 contractual')
    cod = f['Código OT'].strip().upper()
    ton = f['Toneladas metrado']
    if ton <= 0:
        raise ErrorFila('las toneladas deben ser mayores que 0')
    estado = 'ENTREGADO' if f.get('Fecha fin real (acta)') else ('PERDIDO' if f.get('Resultado cotización') == 'Perdida' else 'EN_CURSO')
    datos = dict(cliente=f.get('Cliente'), nombre=f['Nombre del proyecto'], tipo=f['Tipo de estructura'], ton=ton,
                 fuente_ton=f.get('Fuente de tonelaje'), n_piezas=f.get('N° piezas'), pct_planchas=f.get('% planchas'),
                 piezas_por_t=(f['N° piezas'] / ton) if f.get('N° piezas') else None,
                 m2_pint_t=(f['m² de pintura'] / ton) if f.get('m² de pintura') else None,
                 metros_soldadura=f.get('Metros de soldadura'), montaje=1 if f.get('Incluye montaje') == 'Sí' else 0,
                 presupuesto=f.get('Presupuesto sin IGV'), moneda=f.get('Moneda') or 'USD', fecha_pedida=f.get('Fecha pedida por el cliente'),
                 inicio=f['Día 1 contractual'], fecha_comprometida=f.get('Fecha fin contractual') or f.get('Fecha cotizada'),
                 ampliacion_dias=f.get('Ampliación de plazo (días)'), fecha_fin_real=f.get('Fecha fin real (acta)'),
                 penalidad_aplicada=f.get('Penalidad aplicada'), resultado_cot=f.get('Resultado cotización'),
                 plazo_competidor=f.get('Plazo del competidor (días)'), motivo_perdida=f.get('Motivo si se perdió'), estado=estado)
    if datos['fecha_comprometida'] and datos['fecha_comprometida'] < datos['inicio']:
        raise ErrorFila('la fecha fin contractual es anterior al día 1')
    r = cx.execute('SELECT id FROM proyecto_activo WHERE codigo=?', (cod,)).fetchone()
    if r:
        sets = {k: v for k, v in datos.items() if v is not None}
        cx.execute(f'UPDATE proyecto_activo SET {", ".join(k + "=?" for k in sets)} WHERE id=?', (*sets.values(), r[0]))
        return 'actualizada', cod
    cx.execute(f'INSERT INTO proyecto_activo (codigo, creado_por, {", ".join(datos)}) VALUES (?,?,{",".join("?" * len(datos))})',
               (cod, u['id'], *datos.values()))
    return 'nueva', cod


def h_cotiz(cx, f, u):
    _requerir(f, 'Código OT', 'Proceso')
    pid, proc = _ot(cx, f['Código OT']), int(f['Proceso'][:2])
    existe = cx.execute('SELECT 1 FROM cotiz_proceso WHERE proyecto_id=? AND proceso=?', (pid, proc)).fetchone()
    cx.execute('INSERT OR REPLACE INTO cotiz_proceso VALUES (?,?,?,?,?,?,?,?)',
               (pid, proc, f.get('HH estimadas'), f.get('Días estimados'), f.get('Cuadrilla supuesta'), f.get('Contratista previsto'),
                f.get('Estimado por'), f.get('Fecha de estimación')))
    return ('actualizada' if existe else 'nueva'), f['Código OT']


def h_pieza(cx, f, u):
    _requerir(f, 'Código OT', 'N° item', 'Tipo de elemento', 'Cantidad', 'Peso unit. neto (kg)')
    pid = _ot(cx, f['Código OT'])
    err = error_orden_etapas({COL_FECHA_PIEZA[c]: f.get(c) for c in ORDEN_FECHAS_PIEZA})
    if err:
        raise ErrorFila(err)
    pu = f['Peso unit. neto (kg)']
    clase = 'Liviana (0-250 kg)' if pu <= 250 else ('Mediana (251-1000 kg)' if pu <= 1000 else 'Pesada (>1000 kg)')
    conjunto = f.get('Conjunto / marca') or ''
    existe = cx.execute('SELECT 1 FROM pieza WHERE proyecto_id=? AND conjunto=? AND n_item=?', (pid, conjunto, f['N° item'])).fetchone()
    cols = dict(tipo_elemento=f['Tipo de elemento'], perfil=f.get('Perfil'), longitud_mm=f.get('Longitud (mm)'), cantidad=f['Cantidad'],
                peso_unit=pu, peso_total=pu * f['Cantidad'], clase_peso=clase, area_m2=f.get('Área total (m²)'), contratista=f.get('Contratista'),
                orden_fab=f.get('O.F.'), pu_kg=f.get('P.U. (S/ por kg)'), **{COL_FECHA_PIEZA[c]: f.get(c) for c in ORDEN_FECHAS_PIEZA})
    if existe:
        cx.execute(f'UPDATE pieza SET {", ".join(k + "=?" for k in cols)} WHERE proyecto_id=? AND conjunto=? AND n_item=?',
                   (*cols.values(), pid, conjunto, f['N° item']))
        return 'actualizada', f['Código OT']
    cx.execute(f'INSERT INTO pieza (proyecto_id, conjunto, n_item, {", ".join(cols)}) VALUES (?,?,?,{",".join("?" * len(cols))})',
               (pid, conjunto, f['N° item'], *cols.values()))
    return 'nueva', f['Código OT']


def h_tareo(cx, f, u):
    _requerir(f, 'Fecha', 'Código OT', 'Proceso', 'N° personas', 'Horas por persona')
    if f['Fecha'] > date.today().isoformat():
        raise ErrorFila('la fecha del tareo no puede ser futura')
    if f['N° personas'] < 1:
        raise ErrorFila('N° personas debe ser al menos 1')
    pid, proc = _ot(cx, f['Código OT']), int(f['Proceso'][:2])
    contr = None if f.get('Contratista / cuadrilla') in (None, 'Personal propio') else f['Contratista / cuadrilla']
    clave = (f['Fecha'], pid, proc, contr or '', f['N° personas'], f['Horas por persona'])
    if cx.execute('''SELECT 1 FROM tareo WHERE fecha=? AND proyecto_id=? AND proceso=? AND COALESCE(contratista,'')=?
                     AND n_personas=? AND horas_persona=?''', clave).fetchone():
        return 'omitida', f['Código OT']
    cx.execute('''INSERT INTO tareo (fecha, proyecto_id, proceso, contratista, n_personas, horas_persona, kg_avanzados, observacion,
                  registrado_por) VALUES (?,?,?,?,?,?,?,?,?)''',
               (f['Fecha'], pid, proc, contr, f['N° personas'], f['Horas por persona'], f.get('Kg avanzados'),
                f.get('Observación (falta de material, retrabajo, etc.)'), u['id']))
    return 'nueva', f['Código OT']


def h_proceso_fecha(cx, f, u):
    _requerir(f, 'Código OT', 'Proceso', 'Fecha inicio real')
    if f.get('Fecha fin real') and f['Fecha fin real'] < f['Fecha inicio real']:
        raise ErrorFila('la fecha fin es anterior al inicio')
    pid, proc = _ot(cx, f['Código OT']), int(f['Proceso'][:2])
    existe = cx.execute('SELECT 1 FROM proceso_fecha WHERE proyecto_id=? AND proceso=?', (pid, proc)).fetchone()
    cx.execute('INSERT OR REPLACE INTO proceso_fecha VALUES (?,?,?,?,?,?)',
               (pid, proc, f['Fecha inicio real'], f.get('Fecha fin real'), f.get('Cuadrilla promedio'), f.get('Observación')))
    return ('actualizada' if existe else 'nueva'), f['Código OT']


def h_parada(cx, f, u):
    _requerir(f, 'Máquina', 'Fecha y hora inicio', 'Fecha y hora fin', '¿Planificada?', 'Causa')
    if f['Fecha y hora fin'] <= f['Fecha y hora inicio']:
        raise ErrorFila('el fin de la parada debe ser posterior al inicio')
    pid = _ot(cx, f['Código OT afectada']) if f.get('Código OT afectada') else None
    if cx.execute('SELECT 1 FROM parada WHERE maquina=? AND inicio=?', (f['Máquina'], f['Fecha y hora inicio'])).fetchone():
        return 'omitida', f.get('Código OT afectada')
    cx.execute('INSERT INTO parada (maquina, inicio, fin, planificada, causa, proyecto_id, observacion, registrado_por) VALUES (?,?,?,?,?,?,?,?)',
               (f['Máquina'], f['Fecha y hora inicio'], f['Fecha y hora fin'], 1 if f['¿Planificada?'] == 'Sí' else 0, f['Causa'], pid,
                f.get('Observación'), u['id']))
    return 'nueva', f.get('Código OT afectada')


def h_servicio(cx, f, u):
    _requerir(f, 'Código OT', 'Servicio', 'Fecha envío')
    if f.get('Fecha retorno') and f['Fecha retorno'] < f['Fecha envío']:
        raise ErrorFila('el retorno es anterior al envío')
    pid = _ot(cx, f['Código OT'])
    clave = (pid, f['Servicio'], f.get('Proveedor') or '', f['Fecha envío'])
    existe = cx.execute('SELECT id FROM servicio_externo WHERE proyecto_id=? AND servicio=? AND proveedor=? AND fecha_envio=?', clave).fetchone()
    if existe:
        cx.execute('UPDATE servicio_externo SET fecha_retorno=?, kg=?, m2=?, observacion=? WHERE id=?',
                   (f.get('Fecha retorno'), f.get('Kg'), f.get('m²'), f.get('Observación'), existe[0]))
        return 'actualizada', f['Código OT']
    cx.execute('INSERT INTO servicio_externo (proyecto_id, servicio, proveedor, fecha_envio, fecha_retorno, kg, m2, observacion) VALUES (?,?,?,?,?,?,?,?)',
               (*clave, f.get('Fecha retorno'), f.get('Kg'), f.get('m²'), f.get('Observación')))
    return 'nueva', f['Código OT']


def h_compra(cx, f, u):
    _requerir(f, 'Código OT', 'Perfil / material', 'Fecha pedido')
    pid = _ot(cx, f['Código OT'])
    clave = (pid, f.get('N° OC') or '', f['Perfil / material'])
    existe = cx.execute('SELECT id FROM compra WHERE proyecto_id=? AND n_oc=? AND material=?', clave).fetchone()
    vals = (f.get('Proveedor'), f.get('Kg'), f['Fecha pedido'], f.get('Fecha prometida'), f.get('Fecha recepción'), f.get('N° certificado / colada'))
    if existe:
        cx.execute('UPDATE compra SET proveedor=?, kg=?, fecha_pedido=?, fecha_prometida=?, fecha_recepcion=?, certificado=? WHERE id=?', (*vals, existe[0]))
        return 'actualizada', f['Código OT']
    cx.execute('INSERT INTO compra (proyecto_id, n_oc, material, proveedor, kg, fecha_pedido, fecha_prometida, fecha_recepcion, certificado) VALUES (?,?,?,?,?,?,?,?,?)',
               (*clave, *vals))
    return 'nueva', f['Código OT']


def h_evento(cx, f, u):
    _requerir(f, 'Código OT', 'Fecha', 'Tipo de evento')
    pid = _ot(cx, f['Código OT'])
    cur = cx.execute('''INSERT OR IGNORE INTO evento (proyecto_id, fecha, tipo, imputable, dias_impacto, pidio_ampliacion, descripcion)
                        VALUES (?,?,?,?,?,?,?)''', (pid, f['Fecha'], f['Tipo de evento'], f.get('Imputable a'), f.get('Días de impacto'),
                                                   1 if f.get('¿Se pidió ampliación?') == 'Sí' else 0, f.get('Descripción') or ''))
    return ('nueva' if cur.rowcount else 'omitida'), f['Código OT']


MANEJADORES = {'PROYECTO': h_proyecto, 'COTIZ_PROCESO': h_cotiz, 'PIEZAS': h_pieza, 'TAREO': h_tareo, 'PROCESO_FECHAS': h_proceso_fecha,
               'PARADAS': h_parada, 'SERV_EXTERNOS': h_servicio, 'COMPRAS': h_compra, 'EVENTOS': h_evento}


# ------------------------------------------------------------------ importación
def importar(cx, contenido, archivo, usuario, aplicar):
    try:
        wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True, read_only=False)
    except Exception as e:
        raise ValueError(f'No se pudo leer el archivo como Excel (.xlsx): {e}')
    conocidas = [h for h in HOJAS if h[0] in wb.sheetnames]
    if not conocidas:
        raise ValueError('El archivo no tiene ninguna hoja del formato único (PROYECTO, TAREO, PIEZAS, ...)')
    resumen, ots = {}, set()
    cx.execute('SAVEPOINT imp')
    try:
        for nombre, _, _, cols, _ in conocidas:
            ws = wb[nombre]
            titulos = {str(ws.cell(2, j).value).strip(): j for j in range(1, ws.max_column + 1) if ws.cell(2, j).value}
            faltan = [c[0] for c in cols if c[0] not in titulos]
            r = dict(leidas=0, nuevas=0, actualizadas=0, omitidas=0, errores=[])
            resumen[nombre] = r
            if faltan:
                r['errores'].append(dict(fila=2, msg='faltan columnas: ' + ', '.join(faltan)))
                continue
            marcador = ws.cell(3, len(cols) + 1).value
            for i in range(3, ws.max_row + 1):
                crudo = {t: ws.cell(i, titulos[t]).value for t, _, _ in cols}
                if all(v is None or (isinstance(v, str) and not v.strip()) for v in crudo.values()):
                    continue
                if i == 3 and marcador and 'EJEMPLO' in str(marcador).upper():
                    continue
                r['leidas'] += 1
                try:
                    f = {t: convertir(crudo[t], tipo) for t, _, tipo in cols if not (tipo or '').startswith('formula')}
                    estado, ot = MANEJADORES[nombre](cx, f, usuario)
                    r[{'nueva': 'nuevas', 'actualizada': 'actualizadas', 'omitida': 'omitidas'}[estado]] += 1
                    if ot:
                        ots.add(str(ot).upper())
                except ErrorFila as e:
                    r['errores'].append(dict(fila=i, msg=str(e)))
        total_err = sum(len(r['errores']) for r in resumen.values())
        out = dict(archivo=archivo, aplicado=bool(aplicar), hojas=resumen, ots=sorted(ots), errores=total_err,
                   filas=sum(r['leidas'] for r in resumen.values()),
                   cambios=sum(r['nuevas'] + r['actualizadas'] for r in resumen.values()))
    except Exception:
        cx.execute('ROLLBACK TO imp')
        cx.execute('RELEASE imp')
        raise
    if aplicar:
        cx.execute('INSERT INTO importacion (archivo, usuario_id, resumen) VALUES (?,?,?)', (archivo, usuario['id'], json.dumps(out)))
        cx.execute('INSERT INTO actividad (usuario_id, tipo, detalle) VALUES (?,?,?)',
                   (usuario['id'], 'importacion', f'Importó {archivo}: {out["cambios"]} registros, {total_err} filas con error'))
        cx.execute('RELEASE imp')
        cx.commit()
    else:
        cx.execute('ROLLBACK TO imp')
        cx.execute('RELEASE imp')
    return out


def tipo_modelo(tipo):
    """El modelo solo conoce 4 tipos; 'Equipo especial / otro' se cotiza como planta industrial (supuesto conservador)."""
    return tipo if tipo in TIPOS else TIPOS[1]
