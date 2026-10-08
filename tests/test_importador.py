import sqlite3
import sys
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / 'scripts'))
from generar_formato_unico import main as plantilla  # noqa: E402
from plataforma import db_plataforma as dbp  # noqa: E402
from plataforma.avance import ETAPAS, curva_s, fraccion_pieza, resumen_piezas  # noqa: E402
from plataforma.importador import importar  # noqa: E402

U = {'id': 1}
HOY = date.today()


@pytest.fixture
def cx():
    c = sqlite3.connect(':memory:')
    c.row_factory = sqlite3.Row
    c.executescript(dbp.ESQUEMA)
    dbp._migrar(c)
    dbp.crear_usuario(c, 'test', 'Test', 'gerencia', clave='clave-de-prueba')
    c.commit()
    return c


def libro(tmp_path, filas):
    """Plantilla con filas por hoja: {hoja: [{título: valor}]}; borra la fila de ejemplo."""
    wb = openpyxl.load_workbook(plantilla(tmp_path / 'f.xlsx'))
    for hoja, regs in filas.items():
        ws = wb[hoja]
        tit = {str(ws.cell(2, j).value): j for j in range(1, ws.max_column + 1) if ws.cell(2, j).value}
        for j in range(1, ws.max_column + 2):
            if not str(ws.cell(3, j).value or '').startswith('='):
                ws.cell(3, j).value = None
        for i, f in enumerate(regs, start=3):
            for t, v in f.items():
                ws.cell(i, tit[t], v)
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


PROY = {'Código OT': 'ot-t1', 'Nombre del proyecto': 'Prueba', 'Tipo de estructura': 'Nave industrial / packing / cobertura',
        'Toneladas metrado': 10, 'Día 1 contractual': (HOY - timedelta(days=20)).isoformat()}
D = lambda n: (HOY - timedelta(days=n)).isoformat()


def test_validar_no_escribe_y_detecta_errores(cx, tmp_path):
    datos = libro(tmp_path, {
        'PROYECTO': [PROY],
        'TAREO': [{'Fecha': D(3), 'Código OT': 'OT-T1', 'Proceso': '08 Armado', 'N° personas': 4, 'Horas por persona': 8},
                  {'Fecha': (HOY + timedelta(days=3)).isoformat(), 'Código OT': 'OT-T1', 'Proceso': '08 Armado', 'N° personas': 4, 'Horas por persona': 8},
                  {'Fecha': D(2), 'Código OT': 'OT-NO-EXISTE', 'Proceso': '08 Armado', 'N° personas': 4, 'Horas por persona': 8},
                  {'Fecha': D(2), 'Código OT': 'OT-T1', 'Proceso': 'Pintar', 'N° personas': 4, 'Horas por persona': 8},
                  {'Fecha': D(2), 'Código OT': 'OT-T1', 'Proceso': '09 Soldeo', 'N° personas': 4, 'Horas por persona': 15}],
        'PIEZAS': [{'Código OT': 'OT-T1', 'N° item': 1, 'Tipo de elemento': 'VIGA', 'Cantidad': 2, 'Peso unit. neto (kg)': 300,
                    'F. habilitado': D(10), 'F. armado': D(12)}]})
    r = importar(cx, datos, 'f.xlsx', U, aplicar=False)
    errores = {e['fila']: e['msg'] for e in r['hojas']['TAREO']['errores']}
    assert set(errores) == {4, 5, 6, 7}
    assert 'futura' in errores[4] and 'no existe' in errores[5] and 'lista' in errores[6] and 'rango' in errores[7]
    assert 'anterior' in r['hojas']['PIEZAS']['errores'][0]['msg']        # armado antes que habilitado
    assert cx.execute('SELECT COUNT(*) FROM proyecto_activo').fetchone()[0] == 0   # validar no escribe
    assert cx.execute('SELECT COUNT(*) FROM tareo').fetchone()[0] == 0


def test_importar_y_reimportar_es_idempotente(cx, tmp_path):
    datos = libro(tmp_path, {
        'PROYECTO': [PROY],
        'PIEZAS': [{'Código OT': 'OT-T1', 'N° item': i, 'Tipo de elemento': 'COLUMNA', 'Conjunto / marca': f'C-{i}', 'Cantidad': 2,
                    'Peso unit. neto (kg)': 500, 'F. habilitado': D(10), 'F. armado': D(8)} for i in (1, 2, 3)],
        'TAREO': [{'Fecha': D(3), 'Código OT': 'OT-T1', 'Proceso': '08 Armado', 'N° personas': 4, 'Horas por persona': 8}],
        'PARADAS': [{'Máquina': 'Mesa CNC', 'Fecha y hora inicio': f'{D(5)} 08:00', 'Fecha y hora fin': f'{D(5)} 10:00',
                     '¿Planificada?': 'No', 'Causa': 'Falla eléctrica'}]})
    r1 = importar(cx, datos, 'f.xlsx', U, aplicar=True)
    assert r1['errores'] == 0 and r1['hojas']['PIEZAS']['nuevas'] == 3 and r1['ots'] == ['OT-T1']
    r2 = importar(cx, datos, 'f.xlsx', U, aplicar=True)
    assert r2['hojas']['PROYECTO']['actualizadas'] == 1 and r2['hojas']['PIEZAS']['actualizadas'] == 3
    assert r2['hojas']['TAREO']['omitidas'] == 1 and r2['hojas']['PARADAS']['omitidas'] == 1
    assert cx.execute('SELECT COUNT(*) FROM pieza').fetchone()[0] == 3
    assert cx.execute('SELECT COUNT(*) FROM tareo').fetchone()[0] == 1
    assert cx.execute("SELECT peso_total, clase_peso FROM pieza LIMIT 1").fetchone()[:] == (1000.0, 'Mediana (251-1000 kg)')
    assert cx.execute('SELECT COUNT(*) FROM importacion').fetchone()[0] == 2


def test_avance_ponderado_y_curva_s():
    hoy = HOY.isoformat()
    piezas = [dict(peso_total=1000, tipo_elemento='VIGA', contratista='A', clase_peso='Mediana', f_habilitado=D(5), f_armado=D(3),
                   f_soldeo=None, f_liberacion=None, f_granallado=None, f_pintura=None, f_despacho=None),
              dict(peso_total=1000, tipo_elemento='VIGA', contratista='A', clase_peso='Mediana', **{e[0]: D(6) for e in ETAPAS})]
    assert fraccion_pieza(piezas[0], hoy) == pytest.approx(0.35)
    assert fraccion_pieza(piezas[1], hoy) == pytest.approx(1.0)
    r = resumen_piezas(piezas, hoy)
    assert r['kg_avanzado'] == pytest.approx(1350) and r['avance'] == pytest.approx(67.5)
    prog = [(0, 5), (4, 7), (7, 12), (7, 11), (7, 12), (7, 9), (10, 14), (11, 20), (13, 24), (18, 26), (26, 28), (28, 36), (36, 38)]
    c = curva_s(D(20), prog, [10.0] * 13, piezas, [dict(fecha=D(3), hh=32)], hoy, 2000)
    assert c['plan_kg'][-1] == pytest.approx(2000) and c['plan_hh'][-1] == pytest.approx(130)
    reales = [x for x in c['real_kg'] if x is not None]
    assert reales == sorted(reales) and reales[-1] == pytest.approx(1350)
    assert c['real_hh'][-1] is None and max(x for x in c['real_hh'] if x is not None) == pytest.approx(32)
